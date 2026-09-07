# /// script
# requires-python = ">=3.9"
# dependencies = ["openpyxl==3.1.5"]
# ///
"""
nse-sheets-build — bulk-write per-symbol Yahoo chart JSON (from
nse-history-fetch) into one sheet per symbol in a CLOSED xlsx workbook via
openpyxl. Sheet name == ticker symbol exactly, so the Pushpender tab's
HYPERLINK/INDIRECT auto-detect formula lights up the moment a sheet lands.

Cross-platform: pure Python + openpyxl, runs fine on Windows/Linux/macOS.

CRITICAL: the target workbook must be CLOSED in Excel before running this
(openpyxl writing to a file Excel has open corrupts/loses the write, and a
second process fighting Excel over the same file can hang).

Usage:
  uv run --script run.py --workbook <path.xlsx> --raw-dir <dir> [--skip-existing]
"""
import argparse
import json
import os
import re
from datetime import datetime, timezone

from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

HEADER_FILL = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
HEADER_FONT = Font(bold=True, color="FFFFFF")
DATE_FMT = "mm/dd/yyyy"
PRICE_FMT = "₹#,##0.00"
VOL_FMT = "#,##0"
COLS = ["Date", "Open", "High", "Low", "Close", "Adj Close", "Volume"]


def sanitize_sheet_name(name: str) -> str:
    # Excel sheet names: <=31 chars, no : \ / ? * [ ]
    cleaned = re.sub(r"[:\\/\?\*\[\]]", "_", name)
    return cleaned[:31]


def parse_chart_json(raw: dict):
    result = raw["chart"]["result"][0]
    ts = result["timestamp"]
    quote = result["indicators"]["quote"][0]
    adjclose = result["indicators"].get("adjclose", [{}])[0].get("adjclose")
    rows = []
    for i, t in enumerate(ts):
        o, h, l, c, v = (quote["open"][i], quote["high"][i], quote["low"][i],
                          quote["close"][i], quote["volume"][i])
        if o is None or c is None:
            continue  # Yahoo sometimes nulls out a bad session
        date = datetime.fromtimestamp(t, tz=timezone.utc).strftime("%Y-%m-%d")
        rows.append([date, o, h, l, c, (adjclose[i] if adjclose else c), v])
    return rows


def write_symbol_sheet(wb, symbol: str, rows: list):
    sheet_name = sanitize_sheet_name(symbol)
    if sheet_name in wb.sheetnames:
        del wb[sheet_name]
    ws = wb.create_sheet(sheet_name)

    ws.append(COLS)
    for r in range(1, 8):
        cell = ws.cell(row=1, column=r)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL

    for row in rows:
        ws.append(row)

    n = len(rows) + 1  # incl header
    for r in range(2, n + 1):
        ws.cell(row=r, column=1).number_format = DATE_FMT
        for c in range(2, 7):
            ws.cell(row=r, column=c).number_format = PRICE_FMT
        ws.cell(row=r, column=7).number_format = VOL_FMT

    widths = [12, 12, 12, 12, 12, 12, 14]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w

    ws.freeze_panes = "A2"

    if n > 1:
        table_ref = f"A1:G{n}"
        # Table names must be unique workbook-wide and alnum/underscore only
        safe_tbl_name = re.sub(r"[^A-Za-z0-9_]", "_", f"T_{sheet_name}")
        tbl = Table(displayName=safe_tbl_name, ref=table_ref)
        tbl.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium2", showRowStripes=True, showFirstColumn=False,
            showLastColumn=False, showColumnStripes=False,
        )
        ws.add_table(tbl)

    return sheet_name


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--workbook", required=True)
    ap.add_argument("--raw-dir", required=True)
    ap.add_argument("--skip-existing", action="store_true",
                     help="skip symbols that already have a matching sheet")
    args = ap.parse_args()

    wb = load_workbook(args.workbook)

    files = sorted(f for f in os.listdir(args.raw_dir) if f.endswith(".json"))
    total = len(files)
    built = 0
    skipped = 0
    errors = []

    print(f"Building sheets for {total} symbols into {args.workbook}", flush=True)

    for i, fname in enumerate(files, 1):
        symbol = fname[:-5]
        sheet_name = sanitize_sheet_name(symbol)
        if args.skip_existing and sheet_name in wb.sheetnames:
            skipped += 1
            continue
        path = os.path.join(args.raw_dir, fname)
        try:
            with open(path) as f:
                raw = json.load(f)
            rows = parse_chart_json(raw)
            if not rows:
                errors.append((symbol, "no rows"))
                continue
            write_symbol_sheet(wb, symbol, rows)
            built += 1
        except Exception as e:
            errors.append((symbol, str(e)))
        if i % 25 == 0 or i == total:
            print(f"[{i}/{total}] built={built} skipped={skipped} errors={len(errors)}", flush=True)

    wb.save(args.workbook)
    print(f"DONE. built={built} skipped={skipped} errors={len(errors)}", flush=True)
    if errors:
        for sym, err in errors[:20]:
            print(f"  ERROR {sym}: {err}", flush=True)


if __name__ == "__main__":
    main()
